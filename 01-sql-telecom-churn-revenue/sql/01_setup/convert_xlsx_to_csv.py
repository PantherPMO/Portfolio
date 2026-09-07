r"""
============================================================
Project : 01 - Telecommunications Revenue Retention
File    : sql/00_setup/convert_xlsx_to_csv.py
Stage   : PREPARE
Purpose : Convert the six source workbooks to UTF-8 CSV for psql \copy.

THIS IS AN EXTRACTION UTILITY, NOT ANALYSIS - decision D-17.
It touches no values, applies no logic and makes no interpretive choice.
Every join, aggregation, segmentation and KPI calculation in this project is
SQL. See docs/DECISIONS.md D-17.

DEPENDENCY: openpyxl only.  pip install openpyxl
(pandas is deliberately NOT used - openpyxl is the underlying xlsx engine
anyway, so this keeps the extraction step to a single small dependency.)

DO NOT use Excel "Save As CSV" instead. Excel writes DISPLAYED values, which
truncates the high-precision decimals in Total Long Distance Charges and
Total Revenue - that would break reconciliation check REC-08.

Run from the project root:
    python sql\00_setup\convert_xlsx_to_csv.py
============================================================
"""

import csv
import sys
from pathlib import Path

RAW = Path("data/raw")
OUT = Path("data/processed/csv")

FILES = [
    ("Telco_customer_churn_demographics.xlsx", "Telco_Churn", "demographics.csv", 7043, 9),
    ("Telco_customer_churn_location.xlsx",     "Telco_Churn", "location.csv",     7043, 9),
    ("Telco_customer_churn_population.xlsx",   "Population",  "population.csv",   1671, 3),
    ("Telco_customer_churn_services.xlsx",     "Telco_Churn", "services.csv",     7043, 30),
    ("Telco_customer_churn_status.xlsx",       "Telco_Churn", "status.csv",       7043, 11),
    ("Telco_customer_churn.xlsx",              "Telco_Churn", "merged.csv",       7043, 33),
]


def _diagnostics() -> str:
    return (
        f"\n  Interpreter : {sys.executable}"
        f"\n  Version     : {sys.version.split()[0]}"
        f"\n  Working dir : {Path.cwd()}"
    )


def _cell_to_text(value) -> str:
    """Render a cell exactly as stored, at full precision.

    repr-equivalent str() on a float round-trips the stored binary value, so
    the 16-decimal artefacts in Total Long Distance Charges and Total Revenue
    survive intact. The raw layer is all TEXT; casting happens in SQL at the
    promotion step, where a bad value raises rather than coerces silently.
    """
    if value is None:
        return ""
    if isinstance(value, bool):          # not expected, handled for safety
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))           # 150.0 -> "150", not "150.0"
    return str(value)


def main() -> int:
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: openpyxl is not installed for this interpreter." + _diagnostics())
        print(
            "\nInstall it with ONE of:\n"
            f'  "{sys.executable}" -m pip install openpyxl\n'
            "  conda install -c conda-forge openpyxl\n"
            "\nUse the first form if you have several Python installations - it\n"
            "guarantees the package lands in the interpreter this script runs under."
        )
        return 1

    if not RAW.is_dir():
        print(f"ERROR: cannot find {RAW.resolve()}" + _diagnostics())
        print("\nRun this from the project root:")
        print("  cd <path-to-your-clone>/01-sql-telecom-churn-revenue")
        return 1

    OUT.mkdir(parents=True, exist_ok=True)

    failures = []
    print(f"{'file':45s} {'rows':>6s} {'cols':>5s}  status")
    print("-" * 72)

    for xlsx, sheet, csv_name, exp_rows, exp_cols in FILES:
        src = RAW / xlsx
        if not src.exists():
            failures.append(f"MISSING FILE: {xlsx}")
            print(f"{xlsx:45s} {'-':>6s} {'-':>5s}  MISSING")
            continue

        # data_only=True is REQUIRED: these workbooks contain a calculation
        # chain, so without it any formula cell would export as its formula
        # text rather than its value.
        wb = load_workbook(src, read_only=True, data_only=True)
        if sheet not in wb.sheetnames:
            failures.append(f"SHEET NOT FOUND: {xlsx} has no sheet '{sheet}' "
                            f"(found: {', '.join(wb.sheetnames)})")
            print(f"{xlsx:45s} {'-':>6s} {'-':>5s}  BAD SHEET")
            wb.close()
            continue

        ws = wb[sheet]
        rows = cols = 0
        with open(OUT / csv_name, "w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh, lineterminator="\n")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                # read_only mode can emit fully-empty trailing rows
                if all(v is None for v in row):
                    continue
                writer.writerow([_cell_to_text(v) for v in row])
                if i == 0:
                    cols = len(row)
                else:
                    rows += 1
        wb.close()

        ok = (rows == exp_rows and cols == exp_cols)
        if not ok:
            failures.append(
                f"SHAPE: {xlsx} got {rows}x{cols}, expected {exp_rows}x{exp_cols}")
        print(f"{xlsx:45s} {rows:6d} {cols:5d}  {'OK' if ok else 'SHAPE MISMATCH'}")

    print("-" * 72)

    if failures:
        print("\nFAILED - do not proceed to load:")
        for f in failures:
            print("  " + f)
        print(
            "\nA shape mismatch means the source files are not the ones verified at\n"
            "Stage 2. Stop and escalate rather than loading them."
        )
        return 1

    print(f"\nAll six converted to {OUT.resolve()}")
    print("Next: run 00_setup/03_load_raw.sql with psql, from the project root.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
